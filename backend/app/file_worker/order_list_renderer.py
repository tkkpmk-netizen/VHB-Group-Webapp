"""Deterministic Order List XLSX builder and production PDF adapter.

The worker consumes a bounded, immutable JSON snapshot. It never reads the
database and does not evaluate spreadsheet formulas or external references.
"""

import datetime as dt
import hashlib
import json
import re
import subprocess
import zipfile
from decimal import ROUND_HALF_UP, Decimal
from io import BytesIO
from pathlib import Path
from typing import Any, Literal

from openpyxl import Workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.pagebreak import Break
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.worksheet import Worksheet
from pydantic import BaseModel, ConfigDict, Field, field_validator

from app.file_worker.protocol import ArtifactDescriptor

RENDERER_CONTRACT_VERSION = "order-list-spike-v1"
_XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_MONEY_QUANTUM = Decimal("0.01")
_FIXED_ZIP_TIME = (2020, 1, 1, 0, 0, 0)
_FIXED_DOCUMENT_TIMESTAMP = b"2020-01-01T00:00:00Z"


class OrderListLine(BaseModel):
    model_config = ConfigDict(extra="forbid")

    sku: str = Field(min_length=1, max_length=100)
    product_name: str = Field(min_length=1, max_length=500)
    packing: str = Field(min_length=1, max_length=200)
    quantity: Decimal = Field(gt=0, max_digits=18, decimal_places=4)
    unit: str = Field(min_length=1, max_length=30)
    unit_price: Decimal = Field(ge=0, max_digits=18, decimal_places=6)

    @field_validator("sku", "product_name", "packing", "unit")
    @classmethod
    def reject_formula_prefixes(cls, value: str) -> str:
        if value.lstrip().startswith(("=", "+", "-", "@")):
            raise ValueError("Formula-like text is not permitted")
        return value


class OrderListSnapshot(BaseModel):
    model_config = ConfigDict(extra="forbid")

    schema_version: Literal["1.0"] = "1.0"
    order_number: str = Field(min_length=1, max_length=100)
    issue_date: dt.date
    seller_legal_name: str = Field(min_length=1, max_length=300)
    seller_address: str = Field(min_length=1, max_length=1000)
    customer_name: str = Field(min_length=1, max_length=300)
    customer_address: str = Field(min_length=1, max_length=1000)
    currency: str = Field(pattern=r"^[A-Z]{3}$")
    payment_terms: str = Field(min_length=1, max_length=1000)
    incoterm: str = Field(min_length=1, max_length=100)
    destination: str = Field(min_length=1, max_length=300)
    lines: list[OrderListLine] = Field(min_length=1, max_length=1000)
    notes: str | None = Field(default=None, max_length=4000)

    @field_validator(
        "order_number",
        "seller_legal_name",
        "seller_address",
        "customer_name",
        "customer_address",
        "payment_terms",
        "incoterm",
        "destination",
        "notes",
    )
    @classmethod
    def reject_formula_prefixes(cls, value: str | None) -> str | None:
        if value is not None and value.lstrip().startswith(("=", "+", "-", "@")):
            raise ValueError("Formula-like text is not permitted")
        return value


def _money(value: Decimal) -> Decimal:
    return value.quantize(_MONEY_QUANTUM, rounding=ROUND_HALF_UP)


def _sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _artifact(path: Path, content_type: str) -> ArtifactDescriptor:
    data = path.read_bytes()
    return ArtifactDescriptor(
        path=path.name,
        content_type=content_type,
        size_bytes=len(data),
        sha256=_sha256(data),
    )


def _normalize_xlsx_package(path: Path) -> None:
    """Remove ZIP timestamp variance while retaining workbook contents."""
    source = path.read_bytes()
    normalized = BytesIO()
    with zipfile.ZipFile(BytesIO(source)) as current, zipfile.ZipFile(
        normalized, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9
    ) as output:
        for name in sorted(current.namelist()):
            info = zipfile.ZipInfo(name, _FIXED_ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            content = current.read(name)
            if name == "docProps/core.xml":
                content = re.sub(
                    rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*(</dcterms:(?:created|modified)>)",
                    lambda match: match.group(1) + _FIXED_DOCUMENT_TIMESTAMP + match.group(2),
                    content,
                )
            output.writestr(info, content)
    path.write_bytes(normalized.getvalue())


def _build_workbook(snapshot: OrderListSnapshot, output_path: Path) -> Decimal:
    workbook = Workbook()
    workbook.properties.creator = "VHB Group"
    workbook.properties.title = f"Order List {snapshot.order_number}"
    workbook.properties.created = dt.datetime(2020, 1, 1, tzinfo=dt.UTC)
    workbook.properties.modified = dt.datetime(2020, 1, 1, tzinfo=dt.UTC)
    default_sheet = workbook.active
    if default_sheet is not None:
        workbook.remove(default_sheet)
    sheet: Worksheet = workbook.create_sheet("Order List")
    sheet.freeze_panes = "A11"
    sheet.sheet_view.showGridLines = False

    sheet.merge_cells("A1:H1")
    sheet["A1"] = "ORDER LIST / BẢNG ĐẶT HÀNG"
    sheet["A1"].font = Font(name="Liberation Sans", size=18, bold=True, color="17365D")
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 30

    info = (
        ("A3", "Seller / Người bán", "B3", snapshot.seller_legal_name),
        ("A4", "Seller address", "B4", snapshot.seller_address),
        ("A5", "Customer / Khách hàng", "B5", snapshot.customer_name),
        ("A6", "Customer address", "B6", snapshot.customer_address),
        ("F3", "Order No.", "G3", snapshot.order_number),
        ("F4", "Issue date", "G4", snapshot.issue_date.isoformat()),
        ("F5", "Currency", "G5", snapshot.currency),
        ("F6", "Incoterm", "G6", snapshot.incoterm),
        ("F7", "Destination", "G7", snapshot.destination),
    )
    for label_cell, label, value_cell, info_value in info:
        sheet[label_cell] = label
        sheet[label_cell].font = Font(name="Liberation Sans", bold=True, size=9)
        sheet[value_cell] = info_value
        sheet[value_cell].font = Font(name="Liberation Sans", size=9)
    sheet.merge_cells("B3:E3")
    sheet.merge_cells("B4:E4")
    sheet.merge_cells("B5:E5")
    sheet.merge_cells("B6:E6")
    sheet.merge_cells("G3:H3")
    sheet.merge_cells("G4:H4")
    sheet.merge_cells("G5:H5")
    sheet.merge_cells("G6:H6")
    sheet.merge_cells("G7:H7")

    headers = (
        "No.",
        "SKU",
        "Product description",
        "Packing",
        "Quantity",
        "Unit",
        "Unit price",
        "Amount",
    )
    header_row = 10
    thin = Side(style="thin", color="A6A6A6")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for column, value in enumerate(headers, 1):
        cell = sheet.cell(header_row, column, value)
        cell.font = Font(name="Liberation Sans", bold=True, color="FFFFFF", size=9)
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    total = Decimal("0")
    for index, line in enumerate(snapshot.lines, 1):
        row = header_row + index
        amount = _money(line.quantity * line.unit_price)
        total += amount
        values: tuple[Any, ...] = (
            index,
            line.sku,
            line.product_name,
            line.packing,
            line.quantity,
            line.unit,
            line.unit_price,
            amount,
        )
        for column, cell_value in enumerate(values, 1):
            cell = sheet.cell(row, column, cell_value)
            cell.font = Font(name="Liberation Sans", size=8)
            cell.border = border
            cell.alignment = Alignment(
                horizontal="left" if column in {2, 3, 4, 6} else "right",
                vertical="top",
                wrap_text=True,
            )
        sheet.cell(row, 5).number_format = "#,##0.####"
        sheet.cell(row, 7).number_format = '#,##0.000000'
        sheet.cell(row, 8).number_format = '#,##0.00'
        if index % 34 == 0 and index < len(snapshot.lines):
            sheet.row_breaks.append(Break(id=row))

    summary_row = header_row + len(snapshot.lines) + 1
    sheet.merge_cells(start_row=summary_row, start_column=1, end_row=summary_row, end_column=7)
    sheet.cell(summary_row, 1, f"TOTAL / TỔNG CỘNG ({snapshot.currency})")
    sheet.cell(summary_row, 8, _money(total))
    for column in range(1, 9):
        summary_cell = sheet.cell(summary_row, column)
        if not isinstance(summary_cell, MergedCell):
            summary_cell.font = Font(name="Liberation Sans", bold=True, size=9)
            summary_cell.fill = PatternFill("solid", fgColor="D9EAF7")
            summary_cell.border = border
    sheet.cell(summary_row, 8).number_format = '#,##0.00'

    details_row = summary_row + 2
    sheet.merge_cells(start_row=details_row, start_column=1, end_row=details_row, end_column=8)
    sheet.cell(details_row, 1, f"Payment terms / Điều khoản thanh toán: {snapshot.payment_terms}")
    sheet.cell(details_row, 1).alignment = Alignment(wrap_text=True, vertical="top")
    sheet.cell(details_row, 1).font = Font(name="Liberation Sans", size=9)
    if snapshot.notes:
        sheet.merge_cells(
            start_row=details_row + 1,
            start_column=1,
            end_row=details_row + 1,
            end_column=8,
        )
        sheet.cell(details_row + 1, 1, f"Notes / Ghi chú: {snapshot.notes}")
        sheet.cell(details_row + 1, 1).alignment = Alignment(wrap_text=True, vertical="top")
        sheet.cell(details_row + 1, 1).font = Font(name="Liberation Sans", size=9)

    signature_row = details_row + 4
    sheet.merge_cells(start_row=signature_row, start_column=1, end_row=signature_row, end_column=3)
    sheet.merge_cells(start_row=signature_row, start_column=6, end_row=signature_row, end_column=8)
    sheet.cell(signature_row, 1, "SELLER / NGƯỜI BÁN\n(Signature & stamp / Ký và đóng dấu)")
    sheet.cell(signature_row, 6, "BUYER / NGƯỜI MUA\n(Signature / Ký xác nhận)")
    for column in (1, 6):
        sheet.cell(signature_row, column).font = Font(name="Liberation Sans", bold=True, size=9)
        sheet.cell(signature_row, column).alignment = Alignment(
            horizontal="center", vertical="top", wrap_text=True
        )
    sheet.row_dimensions[signature_row].height = 72

    widths = {"A": 7, "B": 15, "C": 38, "D": 22, "E": 13, "F": 10, "G": 15, "H": 17}
    for column_letter, width in widths.items():
        sheet.column_dimensions[column_letter].width = width
    sheet.print_title_rows = f"1:{header_row}"
    sheet.print_area = f"A1:H{signature_row}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    sheet.page_margins.left = 0.25
    sheet.page_margins.right = 0.25
    sheet.page_margins.top = 0.5
    sheet.page_margins.bottom = 0.5
    assert sheet.oddHeader is not None
    assert sheet.oddFooter is not None
    sheet.oddHeader.center.text = f"Order List {snapshot.order_number}"
    sheet.oddFooter.left.text = "VHB Group — Confidential"
    sheet.oddFooter.right.text = "Page &P of &N"

    workbook.save(output_path)
    _normalize_xlsx_package(output_path)
    return _money(total)


def _convert_to_pdf(xlsx_path: Path, output_root: Path, timeout_seconds: int) -> Path:
    command = [
        "libreoffice",
        "--headless",
        "--nologo",
        "--nodefault",
        "--nolockcheck",
        "--nofirststartwizard",
        "-env:UserInstallation=file:///tmp/vhb-lo-profile",
        "--convert-to",
        "pdf:calc_pdf_Export",
        "--outdir",
        str(output_root),
        str(xlsx_path),
    ]
    completed = subprocess.run(
        command,
        check=False,
        capture_output=True,
        timeout=timeout_seconds,
    )
    pdf_path = output_root / f"{xlsx_path.stem}.pdf"
    if completed.returncode != 0 or not pdf_path.is_file():
        raise RuntimeError("LibreOffice could not render the Order List PDF")
    return pdf_path


def render_order_list(
    data: bytes,
    *,
    output_root: Path,
    timeout_seconds: int,
) -> list[ArtifactDescriptor]:
    snapshot = OrderListSnapshot.model_validate_json(data)
    xlsx_path = output_root / "order-list.xlsx"
    total = _build_workbook(snapshot, xlsx_path)
    pdf_path = _convert_to_pdf(xlsx_path, output_root, timeout_seconds)
    metadata = {
        "renderer_contract": RENDERER_CONTRACT_VERSION,
        "schema_version": snapshot.schema_version,
        "order_number": snapshot.order_number,
        "line_count": len(snapshot.lines),
        "currency": snapshot.currency,
        "total": str(total),
        "formula_policy": "values_only",
        "production_renderer": "LibreOffice Calc headless",
        "representative_fixture": True,
        "owner_golden_approval": "pending_real_template_and_100_plus_line_order",
    }
    metadata_path = output_root / "render-metadata.json"
    metadata_path.write_text(
        json.dumps(metadata, ensure_ascii=False, sort_keys=True, separators=(",", ":")),
        encoding="utf-8",
    )
    return [
        _artifact(xlsx_path, _XLSX_CONTENT_TYPE),
        _artifact(pdf_path, "application/pdf"),
        _artifact(metadata_path, "application/json"),
    ]
