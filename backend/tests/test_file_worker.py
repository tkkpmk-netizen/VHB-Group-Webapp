"""T0F manifest, hostile-file, integrity, and lifecycle tests."""

import hashlib
import json
import zipfile
from asyncio import CancelledError, to_thread
from io import BytesIO
from pathlib import Path
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from pydantic import ValidationError

from app.core.config import Settings
from app.file_worker.protocol import (
    FileOperation,
    FileWorkerManifest,
    FileWorkerResult,
    FileWorkerStatus,
    FormulaPolicy,
    InputDescriptor,
    ResourceLimits,
    SecurityPolicy,
)
from app.file_worker.runtime import run
from app.file_worker.scanner import inspect_file
from app.services.file_worker import (
    FileWorkerResultError,
    FileWorkerTimeoutError,
    run_isolated_file_worker,
    validate_file_worker_output,
)


def _sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _manifest(
    data: bytes,
    *,
    filename: str = "input.csv",
    content_type: str = "text/csv",
    operation: FileOperation = FileOperation.tabular_inspect,
    security: SecurityPolicy | None = None,
) -> FileWorkerManifest:
    return FileWorkerManifest(
        job_id=uuid4(),
        operation=operation,
        input=InputDescriptor(
            path="input.bin",
            filename=filename,
            content_type=content_type,
            size_bytes=len(data),
            sha256=_sha(data),
        ),
        security=security or SecurityPolicy(),
    )


def _archive(entries: dict[str, bytes]) -> bytes:
    output = BytesIO()
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, data in entries.items():
            archive.writestr(name, data)
    return output.getvalue()


def _order_list_snapshot(*, line_count: int = 120) -> bytes:
    return json.dumps(
        {
            "schema_version": "1.0",
            "order_number": "OL-2026-0001",
            "issue_date": "2026-08-03",
            "seller_legal_name": "CÔNG TY CỔ PHẦN VIHABA",
            "seller_address": "Hà Nội, Việt Nam",
            "customer_name": "Ava Trading LLC",
            "customer_address": "Dubai, United Arab Emirates",
            "currency": "USD",
            "payment_terms": "30% deposit, 70% before loading",
            "incoterm": "FOB Hai Phong",
            "destination": "Jebel Ali, UAE",
            "notes": "Hàng hóa mới 100%; nhãn phụ theo xác nhận của khách hàng.",
            "lines": [
                {
                    "sku": f"FMCG-{index:04d}",
                    "product_name": f"Sản phẩm thử nghiệm số {index}",
                    "packing": "24 units/carton",
                    "quantity": index,
                    "unit": "CTN",
                    "unit_price": "12.345678",
                }
                for index in range(1, line_count + 1)
            ],
        },
        ensure_ascii=False,
        separators=(",", ":"),
    ).encode()


def test_manifest_rejects_traversal_and_unknown_contract_fields() -> None:
    data = b"a,b\n1,2\n"
    payload = _manifest(data).model_dump(mode="json")
    payload["input"]["path"] = "../escape.csv"
    payload["unexpected"] = True
    with pytest.raises(ValidationError):
        FileWorkerManifest.model_validate(payload)


def test_production_requires_immutable_file_worker_digest() -> None:
    shared = {
        "environment": "production",
        "jwt_secret": "production-secret-that-is-longer-than-32-bytes",
        "storage_secret_key": "production-storage-secret",
    }
    with pytest.raises(ValidationError):
        Settings(**shared, file_worker_image="ghcr.io/example/vhb-file-worker:latest")
    settings = Settings(
        **shared,
        file_worker_image=f"ghcr.io/example/vhb-file-worker@sha256:{'a' * 64}",
    )
    assert "@sha256:" in settings.file_worker_image


def test_runtime_emits_checksummed_result_and_artifact(tmp_path: Path) -> None:
    data = b"name,sku\nCoffee,SKU-1\nTea,SKU-2\n"
    manifest = _manifest(data)
    input_root = tmp_path / "input"
    output_root = tmp_path / "output"
    input_root.mkdir()
    (input_root / "input.bin").write_bytes(data)
    manifest_path = input_root / "manifest.json"
    manifest_path.write_text(manifest.model_dump_json(), encoding="utf-8")

    assert run(manifest_path, input_root, output_root) == 0
    validated = validate_file_worker_output(
        manifest=manifest,
        output_root=output_root,
    )
    assert validated.result.status is FileWorkerStatus.succeeded
    assert validated.result.security_passed is True
    metadata = json.loads(validated.artifacts["tabular-inspection.json"])
    assert metadata["row_count"] == 2
    assert metadata["headers"] == ["name", "sku"]


def test_order_list_renderer_emits_operational_120_line_xlsx_and_pdf(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    data = _order_list_snapshot()
    manifest = _manifest(
        data,
        filename="order-list-snapshot.json",
        content_type="application/json",
        operation=FileOperation.order_list_render,
    )
    input_root = tmp_path / "input"
    output_root = tmp_path / "output"
    input_root.mkdir()
    (input_root / "input.bin").write_bytes(data)
    manifest_path = input_root / "manifest.json"
    manifest_path.write_text(manifest.model_dump_json(), encoding="utf-8")

    def fake_pdf(xlsx_path: Path, output: Path, timeout_seconds: int) -> Path:
        assert xlsx_path.is_file()
        assert timeout_seconds > 0
        path = output / "order-list.pdf"
        path.write_bytes(b"%PDF-1.4\n% representative test artifact\n%%EOF\n")
        return path

    monkeypatch.setattr("app.file_worker.order_list_renderer._convert_to_pdf", fake_pdf)
    assert run(manifest_path, input_root, output_root) == 0
    validated = validate_file_worker_output(manifest=manifest, output_root=output_root)
    assert set(validated.artifacts) == {
        "order-list.xlsx",
        "order-list.pdf",
        "render-metadata.json",
    }

    workbook = load_workbook(BytesIO(validated.artifacts["order-list.xlsx"]), data_only=False)
    sheet = workbook["Order List"]
    assert sheet.max_row == 137
    assert "A1:H1" in {str(item) for item in sheet.merged_cells.ranges}
    assert sheet.print_area == "'Order List'!$A$1:$H$137"
    assert len(sheet.row_breaks.brk) == 3
    assert sheet["C11"].value == "Sản phẩm thử nghiệm số 1"
    assert sheet["H131"].value == pytest.approx(89629.63)
    assert not any(
        cell.data_type == "f"
        for row in sheet.iter_rows()
        for cell in row
        if not isinstance(cell, MergedCell)
    )
    metadata = json.loads(validated.artifacts["render-metadata.json"])
    assert metadata["line_count"] == 120
    assert metadata["formula_policy"] == "values_only"
    assert metadata["owner_golden_approval"].startswith("pending_")

    second_output_root = tmp_path / "second-output"
    assert run(manifest_path, input_root, second_output_root) == 0
    second = validate_file_worker_output(manifest=manifest, output_root=second_output_root)
    assert second.artifacts["order-list.xlsx"] == validated.artifacts["order-list.xlsx"]


def test_order_list_renderer_rejects_formula_injection(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    payload = json.loads(_order_list_snapshot(line_count=1))
    payload["lines"][0]["product_name"] = "=HYPERLINK(\"https://invalid\")"
    data = json.dumps(payload).encode()
    manifest = _manifest(
        data,
        filename="order-list-snapshot.json",
        content_type="application/json",
        operation=FileOperation.order_list_render,
    )
    input_root = tmp_path / "input"
    output_root = tmp_path / "output"
    input_root.mkdir()
    (input_root / "input.bin").write_bytes(data)
    manifest_path = input_root / "manifest.json"
    manifest_path.write_text(manifest.model_dump_json(), encoding="utf-8")
    monkeypatch.setattr(
        "app.file_worker.order_list_renderer._convert_to_pdf",
        lambda *_args, **_kwargs: output_root / "never.pdf",
    )

    assert run(manifest_path, input_root, output_root) == 3
    result = FileWorkerResult.model_validate_json((output_root / "result.json").read_bytes())
    assert result.error_code == "ORDER_LIST_RENDER_FAILED"


@pytest.mark.parametrize(
    ("entries", "code"),
    [
        ({"../escape": b"content"}, "ARCHIVE_PATH_TRAVERSAL"),
        ({"xl/vbaProject.bin": b"macro"}, "MACRO_BLOCKED"),
        ({"xl/externalLinks/externalLink1.xml": b"<external/>"}, "EXTERNAL_LINK_BLOCKED"),
        (
            {
                "xl/worksheets/sheet1.xml": (
                    b'<worksheet xmlns="http://schemas.openxmlformats.org/'
                    b'spreadsheetml/2006/main"><sheetData><row><c><f>1+1</f>'
                    b"</c></row></sheetData></worksheet>"
                )
            },
            "FORMULA_BLOCKED",
        ),
    ],
)
def test_hostile_xlsx_constructs_are_blocked(
    entries: dict[str, bytes],
    code: str,
) -> None:
    data = _archive(entries)
    inspection = inspect_file(
        data,
        filename="hostile.xlsx",
        declared_content_type=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ),
        limits=ResourceLimits(),
        security=SecurityPolicy(formula_policy=FormulaPolicy.reject),
        include_tabular_metadata=False,
    )
    assert inspection.passed is False
    assert code in {finding.code for finding in inspection.findings}


def test_archive_ratio_limit_blocks_zip_bomb_shape() -> None:
    data = _archive({"xl/worksheets/sheet1.xml": b"0" * 2_000_000})
    inspection = inspect_file(
        data,
        filename="bomb.xlsx",
        declared_content_type=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ),
        limits=ResourceLimits(max_archive_ratio=10),
        security=SecurityPolicy(),
        include_tabular_metadata=False,
    )
    assert "ARCHIVE_RATIO_EXCEEDED" in {
        finding.code for finding in inspection.findings
    }


def test_csv_formula_and_binary_content_are_blocked() -> None:
    for data, expected in (
        (b"name,value\nCoffee,=IMPORTDATA(\"https://evil\")\n", "FORMULA_BLOCKED"),
        (b"name,value\nCoffee,\x00binary\n", "TEXT_BINARY_CONTENT"),
    ):
        inspection = inspect_file(
            data,
            filename="input.csv",
            declared_content_type="text/csv",
            limits=ResourceLimits(),
            security=SecurityPolicy(),
            include_tabular_metadata=False,
        )
        assert expected in {finding.code for finding in inspection.findings}


def test_trusted_validator_rejects_tampered_artifact(tmp_path: Path) -> None:
    data = b"name\nCoffee\n"
    manifest = _manifest(data)
    output_root = tmp_path / "output"
    output_root.mkdir()
    artifact = b'{"row_count":1}'
    descriptor = {
        "path": "tabular-inspection.json",
        "content_type": "application/json",
        "size_bytes": len(artifact),
        "sha256": _sha(artifact),
    }
    result = FileWorkerResult(
        job_id=manifest.job_id,
        status=FileWorkerStatus.succeeded,
        input_sha256=manifest.input.sha256,
        security_passed=True,
        detected_content_type="text/csv",
        artifacts=[descriptor],
        tool_versions={
            "file_worker": "1.0.0",
            "protocol": "1.0",
            "python": "3.12",
        },
    )
    (output_root / "result.json").write_text(result.model_dump_json(), encoding="utf-8")
    (output_root / "tabular-inspection.json").write_bytes(artifact + b"tampered")

    with pytest.raises(FileWorkerResultError):
        validate_file_worker_output(manifest=manifest, output_root=output_root)


def test_runtime_rejects_credential_environment(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    data = b"name\nCoffee\n"
    manifest = _manifest(data)
    input_root = tmp_path / "input"
    output_root = tmp_path / "output"
    input_root.mkdir()
    (input_root / "input.bin").write_bytes(data)
    manifest_path = input_root / "manifest.json"
    manifest_path.write_text(manifest.model_dump_json(), encoding="utf-8")
    monkeypatch.setenv("DATABASE_URL", "must-not-enter-file-worker")

    assert run(manifest_path, input_root, output_root) == 3
    result = FileWorkerResult.model_validate_json(
        (output_root / "result.json").read_bytes()
    )
    assert result.error_code == "FORBIDDEN_CREDENTIAL_ENV"
    assert "must-not-enter" not in (output_root / "result.json").read_text()


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("raised", "expected"),
    [
        (TimeoutError(), FileWorkerTimeoutError),
        (CancelledError(), CancelledError),
    ],
)
async def test_timeout_and_cancellation_remove_container_and_temporary_files(
    monkeypatch: pytest.MonkeyPatch,
    raised: BaseException,
    expected: type[BaseException],
) -> None:
    data = b"name\nCoffee\n"
    manifest = _manifest(data)
    commands: list[tuple[str, ...]] = []
    mounted_input: Path | None = None

    class FakeProcess:
        returncode = 0
        killed = False

        async def communicate(self) -> tuple[bytes, bytes]:
            return b"", b""

        def kill(self) -> None:
            self.killed = True

        async def wait(self) -> int:
            return 0

    run_process = FakeProcess()
    cleanup_process = FakeProcess()

    async def fake_subprocess(*args: str, **_: object) -> FakeProcess:
        nonlocal mounted_input
        commands.append(tuple(args))
        if args[:3] == ("docker", "run", "--rm"):
            mount = args[args.index("--mount") + 1]
            source = mount.split("src=", 1)[1].split(",", 1)[0]
            mounted_input = Path(source)
            return run_process
        return cleanup_process

    async def fake_wait_for(*_: object, **__: object) -> tuple[bytes, bytes]:
        raise raised

    monkeypatch.setattr(
        "app.services.file_worker.asyncio.create_subprocess_exec",
        fake_subprocess,
    )
    monkeypatch.setattr("app.services.file_worker.asyncio.wait_for", fake_wait_for)

    with pytest.raises(expected):
        await run_isolated_file_worker(manifest=manifest, input_data=data)

    assert run_process.killed is True
    assert any(command[:3] == ("docker", "rm", "-f") for command in commands)
    run_command = commands[0]
    for required in (
        "--network",
        "--read-only",
        "--cap-drop",
        "--security-opt",
        "--pids-limit",
        "--memory",
        "--cpus",
        "--tmpfs",
    ):
        assert required in run_command
    assert run_command[run_command.index("--network") + 1] == "none"
    assert mounted_input is not None
    assert not await to_thread(mounted_input.exists)
